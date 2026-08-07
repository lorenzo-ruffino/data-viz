#!/opt/homebrew/bin/python3
# -*- coding: utf-8 -*-
"""
Estrazione dati dall'appendice del report ISTAT "Demografia d'impresa"
pubblicato nel 2016 (dati 2009-2014).

File sorgente:
  /Users/lorenzoruffino/Desktop/IMPRESE/Appendice report demografia d'impresa.xlsx
Edizione: "2009-2014"  (prefisso CSV: ed0914_)

Output tidy/long in:
  /Users/lorenzoruffino/Documents/Progetti/data-viz/Demografia imprese/input/raw/

Engine: openpyxl. Interprete: /opt/homebrew/bin/python3
Rieseguibile: sovrascrive i CSV a ogni run.
"""
from pathlib import Path
import csv
import re
import openpyxl
import pandas as pd

# ---------------------------------------------------------------- costanti
SRC = Path("/Users/lorenzoruffino/Desktop/IMPRESE/Appendice report demografia d'impresa.xlsx")
OUTDIR = Path("/Users/lorenzoruffino/Documents/Progetti/data-viz/Demografia imprese/input/raw")
EDIZIONE = "2009-2014"
PREFIX = "ed0914_"
OUTDIR.mkdir(parents=True, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True)


def sheet(name):
    """Match foglio per nome stripped (i nomi possono avere spazi finali)."""
    target = name.strip()
    for w in wb.worksheets:
        if w.title.strip() == target:
            return w
    raise KeyError(f"Foglio non trovato: {name!r}")


def cell(ws, r, c):
    """Valore cella (1-indexed openpyxl)."""
    return ws.cell(row=r, column=c).value


def fmt(v):
    """Formatta un valore per il CSV a precisione piena.
    None -> ''  ; int -> intero ; float intero -> intero ; float -> repr (round-trip)."""
    if v is None:
        return ""
    if isinstance(v, bool):
        return "True" if v else "False"
    if isinstance(v, int):
        return str(v)
    if isinstance(v, float):
        if v.is_integer():
            return str(int(v))
        return repr(v)
    return str(v).strip()


def is_num(v):
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def strip_or_none(v):
    if v is None:
        return None
    s = str(v).strip()
    return s if s != "" else None


YEAR_RE = re.compile(r"(\d{4})")
STRICT_YEAR_RE = re.compile(r"^(\d{4})\s*(\*|\(a\))?$", re.IGNORECASE)


def parse_year(raw):
    """Estrae anno intero e marca provvisorio da un token tipo '2014', '2014(a)', '2009*'."""
    if raw is None:
        return None, False
    s = str(raw).strip()
    m = YEAR_RE.search(s)
    if not m:
        return None, False
    year = int(m.group(1))
    prov = ("*" in s) or ("(a)" in s.lower())
    return year, prov


def strict_year(raw):
    """Come parse_year ma solo se l'INTERA cella e' un token-anno (evita note tipo
    'Stima per le cessate al 2014' che conterrebbero un anno)."""
    if raw is None:
        return None, False
    m = STRICT_YEAR_RE.match(str(raw).strip())
    if not m:
        return None, False
    return int(m.group(1)), bool(m.group(2))


def write_csv(fname, fieldnames, rows):
    path = OUTDIR / fname
    with open(path, "w", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        w.writerow(fieldnames)
        for row in rows:
            w.writerow([fmt(row.get(k)) for k in fieldnames])
    return path


# =====================================================================
# TAVOLA 1 -> ed0914_macrosettori.csv
# Blocchi macrosettore, sotto ciascuno gli anni 2009..2014.
# Colonne file: Anni | Tassi natalita | Imprese nate | Tassi mortalita |
#               Imprese cessate | Tasso netto turnover
# provvisorio = True SOLO per 2014 (nota "Stima per le cessate al 2014").
# =====================================================================
MACRO_T1 = {"Industria in senso stretto", "Costruzioni", "Commercio",
            "Altri servizi", "Totale"}


def extract_tavola1():
    ws = sheet("Tavola 1")
    rows = []
    current = None
    for r in range(1, ws.max_row + 1):
        c1 = strip_or_none(cell(ws, r, 1))
        if c1 is None:
            continue
        if c1 in MACRO_T1:
            current = c1
            continue
        year, _ = strict_year(c1)
        if (year is not None and 2009 <= year <= 2014 and current is not None
                and is_num(cell(ws, r, 2))):
            rows.append({
                "edizione": EDIZIONE,
                "macrosettore": current,
                "anno": year,
                "provvisorio": (year == 2014),
                "tasso_natalita": cell(ws, r, 2),
                "imprese_nate": cell(ws, r, 3),
                "tasso_mortalita": cell(ws, r, 4),
                "imprese_cessate": cell(ws, r, 5),
                "turnover_netto": cell(ws, r, 6),
            })
    return rows


# =====================================================================
# TAVOLA 2 -> ed0914_classi_dipendenti.csv
# Contenuto sfalsato in basso. Blocco "Tassi di natalita" e blocco
# "Tassi di mortalita", 5 classi ciascuno (0, 1-4, 5-9, 10+, Totale).
# =====================================================================
def extract_tavola2():
    ws = sheet("Tavola 2")

    # riga intestazione "Classe di dipendenti" -> gli anni sono nella riga sotto
    header_r = None
    for r in range(1, ws.max_row + 1):
        if strip_or_none(cell(ws, r, 1)) == "Classe di dipendenti":
            header_r = r
            break
    if header_r is None:
        raise RuntimeError("Tavola 2: header 'Classe di dipendenti' non trovato")

    year_r = header_r + 1
    years = []
    for c in range(2, ws.max_column + 1):
        y, _ = parse_year(cell(ws, year_r, c))
        if y is not None:
            years.append((c, y))
    years = years[:6]  # 2009..2014
    assert len(years) == 6, f"Tavola 2: attesi 6 anni, trovati {len(years)}"

    def read_block(label):
        """Trova la riga etichetta (col2 == label) e legge le 5 classi seguenti."""
        lab_r = None
        for r in range(header_r, ws.max_row + 1):
            if strip_or_none(cell(ws, r, 2)) == label:
                lab_r = r
                break
        if lab_r is None:
            raise RuntimeError(f"Tavola 2: etichetta {label!r} non trovata")
        out = {}
        for r in range(lab_r + 1, lab_r + 6):
            classe = strip_or_none(cell(ws, r, 1))
            vals = {}
            for (c, y) in years:
                vals[y] = cell(ws, r, c)
            out[classe] = vals
        return out

    nat = read_block("Tassi di natalità")
    mort = read_block("Tassi di mortalità")

    classi = ["0", "1-4", "5-9", "10+", "Totale"]
    rows = []
    for classe in classi:
        for (_, y) in years:
            rows.append({
                "edizione": EDIZIONE,
                "classe_dipendenti": classe,
                "anno": y,
                "tasso_natalita": nat.get(classe, {}).get(y),
                "tasso_mortalita": mort.get(classe, {}).get(y),
            })
    return rows


# =====================================================================
# TAVOLA 3 -> ed0914_settori_nace.csv
# Settori NACE x anni 2009..2014(a). Ogni anno occupa 4 colonne:
# natalita, mortalita, turnover, (vuota). Anno j -> col (2 + j*4) [openpyxl].
# 2014 provvisorio (nota "(a) Dati provvisori per la mortalita").
# =====================================================================
CODE_RE = re.compile(r"^\s*([A-Za-z]|\d+(?:\s*-\s*\d+)?)\s*-\s*(.+)$", re.DOTALL)


def split_settore(label):
    """Separa codice NACE e nome. 'Totale' -> ('', 'Totale')."""
    s = str(label).strip()
    m = CODE_RE.match(s)
    if not m:
        return "", s
    codice = m.group(1).strip()
    nome = m.group(2).strip()
    return codice, nome


def extract_tavola3():
    ws = sheet("Tavola 3")

    # riga sub-header (col1 inizia con "Settori") ; anni nella riga sopra
    sub_r = None
    for r in range(1, ws.max_row + 1):
        c1 = strip_or_none(cell(ws, r, 1))
        if c1 and c1.startswith("Settori"):
            sub_r = r
            break
    if sub_r is None:
        raise RuntimeError("Tavola 3: sub-header 'Settori...' non trovato")
    year_r = sub_r - 1

    # anni: blocchi da 4 colonne a partire da col2
    year_cols = []
    for j in range(6):
        c = 2 + j * 4
        y, prov = parse_year(cell(ws, year_r, c))
        year_cols.append((c, y, prov))
    assert all(yc[1] is not None for yc in year_cols), f"Tavola 3 anni: {year_cols}"

    rows = []
    for r in range(sub_r + 1, ws.max_row + 1):
        c1 = strip_or_none(cell(ws, r, 1))
        if not c1:
            continue
        if c1.startswith("Fonte") or c1.startswith("(a)"):
            continue
        if not is_num(cell(ws, r, 2)):  # deve avere il primo tasso numerico
            continue
        codice, nome = split_settore(c1)
        for (c, y, prov_hdr) in year_cols:
            rows.append({
                "edizione": EDIZIONE,
                "settore_codice": codice,
                "settore_nome": nome,
                "anno": y,
                "provvisorio": (y == 2014),  # nota: (a) prov. per la mortalita 2014
                "tasso_natalita": cell(ws, r, c),
                "tasso_mortalita": cell(ws, r, c + 1),
                "turnover_netto": cell(ws, r, c + 2),
            })
    return rows


# =====================================================================
# TAVOLA 4 -> ed0914_regioni.csv
# Regioni/ripartizioni x anni 2009..2014, 3 colonne per anno
# (natalita, mortalita, turnover). Anno j -> col (2 + j*3) [openpyxl].
# 2014 provvisorio (nota "(a) Stima per le cessate al 2014").
# =====================================================================
def extract_tavola4():
    ws = sheet("Tavola 4")

    # riga sub-header con i tre "tasso ..." ; anni nella riga sopra (col1='Aree geografiche')
    sub_r = None
    for r in range(1, ws.max_row + 1):
        if strip_or_none(cell(ws, r, 2)) == "tasso di natalità":
            sub_r = r
            break
    if sub_r is None:
        raise RuntimeError("Tavola 4: sub-header 'tasso di natalità' non trovato")
    year_r = sub_r - 1

    year_cols = []
    for j in range(6):
        c = 2 + j * 3
        y, _ = parse_year(cell(ws, year_r, c))
        year_cols.append((c, y))
    assert all(yc[1] is not None for yc in year_cols), f"Tavola 4 anni: {year_cols}"

    rows = []
    for r in range(sub_r + 1, ws.max_row + 1):
        terr = strip_or_none(cell(ws, r, 1))
        if not terr:
            continue
        if terr.startswith("Fonte") or terr.startswith("(a)"):
            continue
        if not is_num(cell(ws, r, 2)):
            continue
        for (c, y) in year_cols:
            rows.append({
                "edizione": EDIZIONE,
                "territorio": terr,
                "anno": y,
                "provvisorio": (y == 2014),
                "tasso_natalita": cell(ws, r, c),
                "tasso_mortalita": cell(ws, r, c + 1),
                "turnover_netto": cell(ws, r, c + 2),
            })
    return rows


# =====================================================================
# TAVOLA 5 -> ed0914_sopravvivenza.csv
# Matrice triangolare: coorti (anno di nascita) 2009..2013,
# anni di osservazione 2010..2014. Blocchi per macrosettore.
# col1 = macrosettore (solo prima riga blocco), col2 = coorte,
# col3..col7 = anni osservazione. Solo celle non vuote.
# =====================================================================
def extract_tavola5():
    ws = sheet("Tavola 5")

    # riga intestazione: col1 == "Macrosettori", col2 == "anno di nascita"
    hdr_r = None
    for r in range(1, ws.max_row + 1):
        if strip_or_none(cell(ws, r, 1)) == "Macrosettori":
            hdr_r = r
            break
    if hdr_r is None:
        raise RuntimeError("Tavola 5: header 'Macrosettori' non trovato")

    # anni di osservazione da col3 in poi sulla riga header
    obs_cols = []
    for c in range(3, ws.max_column + 1):
        y, _ = parse_year(cell(ws, hdr_r, c))
        if y is not None:
            obs_cols.append((c, y))
    obs_cols = obs_cols[:5]  # 2010..2014
    assert len(obs_cols) == 5, f"Tavola 5: attesi 5 anni oss, trovati {len(obs_cols)}"

    rows = []
    current = None
    for r in range(hdr_r + 1, ws.max_row + 1):
        c1 = strip_or_none(cell(ws, r, 1))
        coorte_raw = cell(ws, r, 2)
        if c1:
            current = c1
        cy, _ = parse_year(coorte_raw)
        if cy is None or not (2009 <= cy <= 2013):
            continue
        if current is None:
            continue
        for (c, oy) in obs_cols:
            v = cell(ws, r, c)
            if v is None or (isinstance(v, str) and v.strip() == ""):
                continue  # cella vuota: la si salta (matrice triangolare)
            rows.append({
                "edizione": EDIZIONE,
                "macrosettore": current,
                "coorte": cy,
                "anno_osservazione": oy,
                "anni_dalla_nascita": oy - cy,
                "tasso_sopravvivenza": v,
            })
    return rows


# =====================================================================
# TAVOLA 6 -> ed0914_addetti_coorte2010.csv
# Addetti della COORTE 2010 (non 2009). Finestra 2010->2014 = 4 anni.
# col2=(a) nate 2010, col3=(b) sopravviventi @2010, col4=(c) sopravviventi @2014,
# col5=(b-a)/a*100, col6=(c-b)/b*100, col7=(c-a)/a*100.
# =====================================================================
MACRO_T6 = {"Industria in senso stretto", "Costruzioni", "Commercio",
            "Altri servizi", "Totale"}


def extract_tavola6():
    ws = sheet("Tavola 6")  # nome reale 'Tavola 6 ' con spazio finale
    rows = []
    for r in range(1, ws.max_row + 1):
        macro = strip_or_none(cell(ws, r, 1))
        if macro not in MACRO_T6:
            continue
        if not is_num(cell(ws, r, 2)):
            continue
        rows.append({
            "edizione": EDIZIONE,
            "coorte": 2010,
            "macrosettore": macro,
            "addetti_t0_nate": cell(ws, r, 2),
            "addetti_t0_sopravviventi": cell(ws, r, 3),
            "addetti_tfin_sopravviventi": cell(ws, r, 4),
            "perdita_pct_da_cessazioni": cell(ws, r, 5),
            "crescita_pct_sopravviventi": cell(ws, r, 6),
            "variazione_pct_netta": cell(ws, r, 7),
            "orizzonte_anni": 4,
        })
    return rows


# =====================================================================
# ESECUZIONE + SCRITTURA
# =====================================================================
specs = [
    ("ed0914_macrosettori.csv",
     ["edizione", "macrosettore", "anno", "provvisorio", "tasso_natalita",
      "imprese_nate", "tasso_mortalita", "imprese_cessate", "turnover_netto"],
     extract_tavola1),
    ("ed0914_classi_dipendenti.csv",
     ["edizione", "classe_dipendenti", "anno", "tasso_natalita", "tasso_mortalita"],
     extract_tavola2),
    ("ed0914_settori_nace.csv",
     ["edizione", "settore_codice", "settore_nome", "anno", "provvisorio",
      "tasso_natalita", "tasso_mortalita", "turnover_netto"],
     extract_tavola3),
    ("ed0914_regioni.csv",
     ["edizione", "territorio", "anno", "provvisorio",
      "tasso_natalita", "tasso_mortalita", "turnover_netto"],
     extract_tavola4),
    ("ed0914_sopravvivenza.csv",
     ["edizione", "macrosettore", "coorte", "anno_osservazione",
      "anni_dalla_nascita", "tasso_sopravvivenza"],
     extract_tavola5),
    ("ed0914_addetti_coorte2010.csv",
     ["edizione", "coorte", "macrosettore", "addetti_t0_nate",
      "addetti_t0_sopravviventi", "addetti_tfin_sopravviventi",
      "perdita_pct_da_cessazioni", "crescita_pct_sopravviventi",
      "variazione_pct_netta", "orizzonte_anni"],
     extract_tavola6),
]

written = {}
for fname, fields, fn in specs:
    data = fn()
    path = write_csv(fname, fields, data)
    written[fname] = (path, len(data))
    print(f"SCRITTO {fname}: {len(data)} righe -> {path}")


# =====================================================================
# VERIFICA: rilettura con pandas + valori di controllo
# =====================================================================
def load(fname):
    return pd.read_csv(written[fname][0])


print("\n================ VERIFICA ================")
ok = True


def check(desc, got, exp, tol=1e-9):
    global ok
    if isinstance(exp, str):
        good = (str(got) == exp)
    else:
        good = (got is not None) and (abs(float(got) - float(exp)) <= tol)
    ok = ok and good
    print(f"[{'OK' if good else 'FAIL'}] {desc}: got={got!r} exp={exp!r}")


# --- macrosettori
m = load("ed0914_macrosettori.csv")
print("ed0914_macrosettori shape:", m.shape)
check("macro n_righe", m.shape[0], 30)          # 5 macrosettori x 6 anni
r = m[(m.macrosettore == "Industria in senso stretto") & (m.anno == 2009)].iloc[0]
check("macro Industria 2009 tasso_natalita", r.tasso_natalita, 4.53676308669081)
check("macro Industria 2009 imprese_nate", r.imprese_nate, 20808)
check("macro Industria 2009 imprese_cessate", r.imprese_cessate, 30935)
r = m[(m.macrosettore == "Totale") & (m.anno == 2014)].iloc[0]
check("macro Totale 2014 tasso_natalita", r.tasso_natalita, 7.135776787076785)
check("macro provvisorio 2014 True", bool(m[m.anno == 2014].provvisorio.all()), True)
check("macro provvisorio <2014 False", bool((~m[m.anno < 2014].provvisorio).all()), True)

# --- classi dipendenti
c = load("ed0914_classi_dipendenti.csv")
print("ed0914_classi_dipendenti shape:", c.shape)
check("classi n_righe", c.shape[0], 30)          # 5 classi x 6 anni
r = c[(c.classe_dipendenti == "0") & (c.anno == 2009)].iloc[0]
check("classi cl0 2009 tasso_natalita", r.tasso_natalita, 8.996909206892981)
check("classi cl0 2009 tasso_mortalita", r.tasso_mortalita, 9.393217649397853)
r = c[(c.classe_dipendenti == "10+") & (c.anno == 2014)].iloc[0]
check("classi cl10+ 2014 tasso_natalita", r.tasso_natalita, 1.191385073362888)

# --- settori nace
s = load("ed0914_settori_nace.csv")
print("ed0914_settori_nace shape:", s.shape)
check("settori righe", s.shape[0], 180)
print("  codici distinti:", sorted(s.settore_codice.dropna().astype(str).unique().tolist()))

# --- regioni
g = load("ed0914_regioni.csv")
print("ed0914_regioni shape:", g.shape)
r = g[(g.territorio == "Piemonte") & (g.anno == 2009)].iloc[0]
check("regioni Piemonte 2009 tasso_natalita", r.tasso_natalita, 7.2)
r = g[(g.territorio == "Italia") & (g.anno == 2009)].iloc[0]
check("regioni Italia 2009 tasso_natalita", r.tasso_natalita, 7.2)
check("regioni n_territori", g.territorio.nunique(), 26)

# --- sopravvivenza
sv = load("ed0914_sopravvivenza.csv")
print("ed0914_sopravvivenza shape:", sv.shape)
check("sopravv n_righe", sv.shape[0], 75)         # 5 macro x (5+4+3+2+1)
r = sv[(sv.macrosettore == "Industria in s.s.") & (sv.coorte == 2009) &
       (sv.anno_osservazione == 2014)].iloc[0]
check("sopravv Industria coorte2009 oss2014", r.tasso_sopravvivenza, 51.66282199154172)
check("sopravv anni_dalla_nascita coerente",
      bool(((sv.anni_dalla_nascita == sv.anno_osservazione - sv.coorte)).all()), True)

# --- addetti coorte 2010
a = load("ed0914_addetti_coorte2010.csv")
print("ed0914_addetti_coorte2010 shape:", a.shape)
check("addetti n_righe", a.shape[0], 5)            # 5 macrosettori
r = a[a.macrosettore == "Industria in senso stretto"].iloc[0]
check("addetti Industria a", r.addetti_t0_nate, 41908)
check("addetti Industria b", r.addetti_t0_sopravviventi, 25950)
check("addetti Industria c", r.addetti_tfin_sopravviventi, 52044.97)
check("addetti coorte==2010", bool((a.coorte == 2010).all()), True)
check("addetti orizzonte==4", bool((a.orizzonte_anni == 4).all()), True)

print("\n==== RISULTATO:", "TUTTI OK" if ok else "CI SONO FAIL", "====")
if not ok:
    raise SystemExit(1)
