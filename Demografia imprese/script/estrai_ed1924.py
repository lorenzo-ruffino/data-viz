#!/opt/homebrew/bin/python3
# -*- coding: utf-8 -*-
"""
Estrazione tavole statistiche dal report ISTAT "Demografia d'impresa".
Edizione: 2019-2024 (2024 provvisorio).

Sorgente : /Users/lorenzoruffino/Desktop/IMPRESE/TAVOLE-STATISTICHE.xlsx
Output   : Demografia imprese/input/raw/ed1924_*.csv  (tidy/long, UTF-8, sep=',', dec='.')
Engine   : openpyxl (data_only=True)

Rieseguibile: sovrascrive i CSV ad ogni run.
"""

import os
import re
import math
import openpyxl
import pandas as pd

SRC = "/Users/lorenzoruffino/Desktop/IMPRESE/TAVOLE-STATISTICHE.xlsx"
OUTDIR = "/Users/lorenzoruffino/Documents/Progetti/data-viz/Demografia imprese/input/raw"
EDIZIONE = "2019-2024"
PREFIX = "ed1924"

os.makedirs(OUTDIR, exist_ok=True)

wb = openpyxl.load_workbook(SRC, data_only=True)


def get_sheet(name_stripped):
    """Trova un foglio matchando il nome ripulito (i nomi possono avere spazi finali)."""
    for s in wb.sheetnames:
        if s.strip() == name_stripped:
            return wb[s]
    raise KeyError(f"Foglio non trovato: {name_stripped!r}. Disponibili: {wb.sheetnames}")


def parse_year(val):
    """Ritorna (anno_int, provvisorio_bool) oppure (None, None) se non e' un anno."""
    if val is None:
        return None, None
    if isinstance(val, bool):
        return None, None
    if isinstance(val, (int, float)):
        f = float(val)
        if f.is_integer() and 1900 <= f <= 2100:
            return int(f), False
        return None, None
    s = str(val).strip()
    prov = ("*" in s) or ("(a)" in s.lower())
    # La cella e' un anno solo se, tolti i marcatori (a)/*, resta SOLO un anno a 4 cifre
    # (evita di catturare "2024" dentro titoli/note tipo "... al 2024." o "Anni 2019-2024").
    core = s.lower().replace("(a)", "").replace("*", "").strip()
    if re.fullmatch(r"(19|20)\d{2}", core):
        return int(core), prov
    return None, None


def clean(val):
    """Strip degli spazi sulle etichette testuali; altri valori invariati."""
    if isinstance(val, str):
        return val.strip()
    return val


def to_int64(df, cols):
    for c in cols:
        # Int64 nullable: preserva interi e celle vuote
        df[c] = df[c].astype("Int64")
    return df


anomalies = []
written = []


def write_csv(df, fname, cols):
    df = df[cols].copy()
    path = os.path.join(OUTDIR, fname)
    df.to_csv(path, index=False, encoding="utf-8")
    return path, df


# ---------------------------------------------------------------------------
# TAVOLA 1 -> ed1924_totale_annuale.csv
# ---------------------------------------------------------------------------
def estrai_tavola1():
    ws = get_sheet("Tavola 1")
    rows = []
    for r in range(1, ws.max_row + 1):
        a = ws.cell(r, 1).value
        anno, prov = parse_year(a)
        if anno is None:
            continue
        rows.append({
            "edizione": EDIZIONE,
            "anno": anno,
            "provvisorio": prov,
            "tasso_natalita": ws.cell(r, 2).value,
            "imprese_nate": ws.cell(r, 3).value,
            "tasso_mortalita": ws.cell(r, 4).value,
            "imprese_cessate": ws.cell(r, 5).value,
            "turnover_netto": ws.cell(r, 6).value,
        })
    df = pd.DataFrame(rows)
    df = to_int64(df, ["imprese_nate", "imprese_cessate"])
    cols = ["edizione", "anno", "provvisorio", "tasso_natalita", "imprese_nate",
            "tasso_mortalita", "imprese_cessate", "turnover_netto"]
    return write_csv(df, f"{PREFIX}_totale_annuale.csv", cols)


# ---------------------------------------------------------------------------
# TAVOLA 2 -> ed1924_macrosettori.csv
# ---------------------------------------------------------------------------
def estrai_tavola2():
    ws = get_sheet("Tavola 2")
    rows = []
    current = None
    for r in range(4, ws.max_row + 1):
        a = ws.cell(r, 1).value
        anno, prov = parse_year(a)
        if anno is not None:
            if current is None:
                anomalies.append(f"Tavola 2: riga {r} anno {anno} senza macrosettore corrente")
                continue
            rows.append({
                "edizione": EDIZIONE,
                "macrosettore": current,
                "anno": anno,
                "provvisorio": prov,
                "tasso_natalita": ws.cell(r, 2).value,
                "imprese_nate": ws.cell(r, 3).value,
                "tasso_mortalita": ws.cell(r, 4).value,
                "imprese_cessate": ws.cell(r, 5).value,
                "turnover_netto": ws.cell(r, 6).value,
            })
        elif isinstance(a, str):
            s = a.strip()
            if s and not s.startswith("("):
                current = s
    df = pd.DataFrame(rows)
    df = to_int64(df, ["imprese_nate", "imprese_cessate"])
    cols = ["edizione", "macrosettore", "anno", "provvisorio", "tasso_natalita",
            "imprese_nate", "tasso_mortalita", "imprese_cessate", "turnover_netto"]
    return write_csv(df, f"{PREFIX}_macrosettori.csv", cols)


# ---------------------------------------------------------------------------
# TAVOLA 3 -> ed1924_settori_tecnologia.csv
#   header anni su r3 (ogni 4 colonne); sotto-intestazioni su r4
# ---------------------------------------------------------------------------
SUBMAP_T3 = {
    "tasso di natalità": "tasso_natalita",
    "imprese nate": "imprese_nate",
    "tasso di mortalità": "tasso_mortalita",
    "imprese cessate": "imprese_cessate",
}


def estrai_tavola3():
    ws = get_sheet("Tavola 3")
    HDR_YEAR = 3
    HDR_SUB = 4
    DATA_START = 5

    # Individua i blocchi anno su r3
    year_blocks = []  # (col_start, anno, prov)
    for c in range(2, ws.max_column + 1):
        anno, prov = parse_year(ws.cell(HDR_YEAR, c).value)
        if anno is not None:
            year_blocks.append((c, anno, prov))

    rows = []
    for r in range(DATA_START, ws.max_row + 1):
        settore = ws.cell(r, 1).value
        if not isinstance(settore, str):
            continue
        settore = settore.strip()
        if not settore or settore.startswith("("):
            continue
        for (c0, anno, prov) in year_blocks:
            rec = {"tasso_natalita": None, "imprese_nate": None,
                   "tasso_mortalita": None, "imprese_cessate": None}
            for off in range(4):
                c = c0 + off
                sub = ws.cell(HDR_SUB, c).value
                if not isinstance(sub, str):
                    continue
                key = SUBMAP_T3.get(sub.strip().lower())
                if key is None:
                    anomalies.append(f"Tavola 3: sotto-intestazione non mappata {sub!r} a col {c}")
                    continue
                rec[key] = ws.cell(r, c).value
            rows.append({
                "edizione": EDIZIONE,
                "settore": settore,
                "anno": anno,
                "provvisorio": prov,
                **rec,
            })
    df = pd.DataFrame(rows)
    df = to_int64(df, ["imprese_nate", "imprese_cessate"])
    cols = ["edizione", "settore", "anno", "provvisorio", "tasso_natalita",
            "imprese_nate", "tasso_mortalita", "imprese_cessate"]
    return write_csv(df, f"{PREFIX}_settori_tecnologia.csv", cols)


# ---------------------------------------------------------------------------
# TAVOLA 4 -> ed1924_regioni.csv
#   header anni su r3 (ogni 3 colonne); sotto-intestazioni su r4
# ---------------------------------------------------------------------------
SUBMAP_T4 = {
    "tasso di natalità": "tasso_natalita",
    "tasso di mortalità": "tasso_mortalita",
    "turnover netto": "turnover_netto",
}


def estrai_tavola4():
    ws = get_sheet("Tavola 4")
    HDR_YEAR = 3
    HDR_SUB = 4
    DATA_START = 5

    year_blocks = []
    for c in range(2, ws.max_column + 1):
        anno, prov = parse_year(ws.cell(HDR_YEAR, c).value)
        if anno is not None:
            year_blocks.append((c, anno, prov))

    rows = []
    for r in range(DATA_START, ws.max_row + 1):
        terr = ws.cell(r, 1).value
        if not isinstance(terr, str):
            continue
        terr = terr.strip()
        if not terr or terr.startswith("("):
            continue
        for (c0, anno, prov) in year_blocks:
            rec = {"tasso_natalita": None, "tasso_mortalita": None, "turnover_netto": None}
            for off in range(3):
                c = c0 + off
                sub = ws.cell(HDR_SUB, c).value
                if not isinstance(sub, str):
                    continue
                key = SUBMAP_T4.get(sub.strip().lower())
                if key is None:
                    anomalies.append(f"Tavola 4: sotto-intestazione non mappata {sub!r} a col {c}")
                    continue
                rec[key] = ws.cell(r, c).value
            rows.append({
                "edizione": EDIZIONE,
                "territorio": terr,
                "anno": anno,
                "provvisorio": prov,
                **rec,
            })
    df = pd.DataFrame(rows)
    cols = ["edizione", "territorio", "anno", "provvisorio",
            "tasso_natalita", "tasso_mortalita", "turnover_netto"]
    return write_csv(df, f"{PREFIX}_regioni.csv", cols)


# ---------------------------------------------------------------------------
# TAVOLA 5 -> ed1924_sopravvivenza.csv
#   matrice triangolare: coorte (col B, anno di nascita) x anno osservazione (r3, col C..G)
# ---------------------------------------------------------------------------
def estrai_tavola5():
    ws = get_sheet("Tavola 5")
    HDR = 3
    # anni di osservazione: colonne dalla C (3) in poi con un anno su r3
    obs_cols = []  # (col, anno)
    for c in range(3, ws.max_column + 1):
        anno, _ = parse_year(ws.cell(HDR, c).value)
        if anno is not None:
            obs_cols.append((c, anno))

    rows = []
    current = None
    for r in range(HDR + 1, ws.max_row + 1):
        b = ws.cell(r, 2).value
        coorte, _ = parse_year(b)
        if coorte is None:
            continue
        a = ws.cell(r, 1).value
        if isinstance(a, str) and a.strip() and not a.strip().startswith("("):
            current = a.strip()
        if current is None:
            anomalies.append(f"Tavola 5: riga {r} coorte {coorte} senza macrosettore")
            continue
        for (c, anno_oss) in obs_cols:
            val = ws.cell(r, c).value
            if val is None:
                continue
            rows.append({
                "edizione": EDIZIONE,
                "macrosettore": current,
                "coorte": coorte,
                "anno_osservazione": anno_oss,
                "anni_dalla_nascita": anno_oss - coorte,
                "tasso_sopravvivenza": val,
            })
    df = pd.DataFrame(rows)
    cols = ["edizione", "macrosettore", "coorte", "anno_osservazione",
            "anni_dalla_nascita", "tasso_sopravvivenza"]
    return write_csv(df, f"{PREFIX}_sopravvivenza.csv", cols)


# ---------------------------------------------------------------------------
# TAVOLA 6 -> ed1924_addetti_coorte2019.csv
# ---------------------------------------------------------------------------
def estrai_tavola6():
    ws = get_sheet("Tavola 6")
    DATA_START = 5
    rows = []
    for r in range(DATA_START, ws.max_row + 1):
        m = ws.cell(r, 1).value
        if not isinstance(m, str):
            continue
        m = m.strip()
        if not m or m.startswith("("):
            continue
        rows.append({
            "edizione": EDIZIONE,
            "coorte": 2019,
            "macrosettore": m,
            "addetti_t0_nate": ws.cell(r, 2).value,
            "addetti_t0_sopravviventi": ws.cell(r, 3).value,
            "addetti_t5_sopravviventi": ws.cell(r, 4).value,
            "perdita_pct_da_cessazioni": ws.cell(r, 5).value,
            "crescita_pct_sopravviventi": ws.cell(r, 6).value,
            "variazione_pct_netta": ws.cell(r, 7).value,
        })
    df = pd.DataFrame(rows)
    cols = ["edizione", "coorte", "macrosettore", "addetti_t0_nate",
            "addetti_t0_sopravviventi", "addetti_t5_sopravviventi",
            "perdita_pct_da_cessazioni", "crescita_pct_sopravviventi",
            "variazione_pct_netta"]
    return write_csv(df, f"{PREFIX}_addetti_coorte2019.csv", cols)


# ---------------------------------------------------------------------------
# GRAFICO WEB -> ed1924_serie_tassi_2006_2024.csv
#   usa il blocco in PERCENTUALE: col E (anno), col F (natalita), col G (mortalita)
# ---------------------------------------------------------------------------
def estrai_grafico_web():
    ws = get_sheet("grafico web")
    COL_ANNO, COL_NAT, COL_MOR = 5, 6, 7  # E, F, G
    rows = []
    for r in range(1, ws.max_row + 1):
        anno, prov = parse_year(ws.cell(r, COL_ANNO).value)
        if anno is None:
            continue
        rows.append({
            "edizione": EDIZIONE,
            "anno": anno,
            "provvisorio": prov,
            "tasso_natalita": ws.cell(r, COL_NAT).value,
            "tasso_mortalita": ws.cell(r, COL_MOR).value,
        })
    df = pd.DataFrame(rows)
    cols = ["edizione", "anno", "provvisorio", "tasso_natalita", "tasso_mortalita"]
    return write_csv(df, f"{PREFIX}_serie_tassi_2006_2024.csv", cols)


# ---------------------------------------------------------------------------
# RUN + VALIDAZIONE
# ---------------------------------------------------------------------------
def approx(a, b, tol=1e-9):
    if a is None or b is None:
        return False
    return math.isclose(float(a), float(b), rel_tol=0.0, abs_tol=tol)


results = {}
for fn in (estrai_tavola1, estrai_tavola2, estrai_tavola3, estrai_tavola4,
           estrai_tavola5, estrai_tavola6, estrai_grafico_web):
    path, df = fn()
    results[os.path.basename(path)] = (path, df)

print("=" * 70)
print("CSV SCRITTI + shape")
print("=" * 70)
for name, (path, df) in results.items():
    print(f"{name:40s} shape={df.shape}  -> {path}")

# --- Controlli ---
print("\n" + "=" * 70)
print("VALIDAZIONE (rilettura dai CSV)")
print("=" * 70)

checks = []


def load(name):
    return pd.read_csv(os.path.join(OUTDIR, name))


def check(desc, cond):
    checks.append((desc, bool(cond)))
    print(f"  [{'OK ' if cond else 'FAIL'}] {desc}")


# ed1924_totale_annuale
d = load(f"{PREFIX}_totale_annuale.csv")
r24 = d[d.anno == 2024].iloc[0]
check("totale_annuale 2024 imprese_nate == 292787", int(r24.imprese_nate) == 292787)
check("totale_annuale 2024 imprese_cessate == 255667", int(r24.imprese_cessate) == 255667)
check("totale_annuale 2024 provvisorio == True", bool(r24.provvisorio) == True)
check("totale_annuale 2021 turnover_netto == 1.1",
      approx(d[d.anno == 2021].iloc[0].turnover_netto, 1.1))
check("totale_annuale 2020 imprese_nate == 245922",
      int(d[d.anno == 2020].iloc[0].imprese_nate) == 245922)

# ed1924_macrosettori
d = load(f"{PREFIX}_macrosettori.csv")
row = d[(d.macrosettore == "Costruzioni") & (d.anno == 2021)].iloc[0]
check("macrosettori Costruzioni 2021 tasso_natalita == 9.7", approx(row.tasso_natalita, 9.7))
row = d[(d.macrosettore == "Industria in senso stretto") & (d.anno == 2020)].iloc[0]
check("macrosettori Industria in s.s. 2020 tasso_natalita == 3.8", approx(row.tasso_natalita, 3.8))
check("macrosettori: 5 macrosettori", d.macrosettore.nunique() == 5)
check("macrosettori: 30 righe (5x6)", len(d) == 30)

# ed1924_settori_tecnologia
d = load(f"{PREFIX}_settori_tecnologia.csv")
row = d[(d.settore == "Servizi tecnologici ad Alto contenuto di conoscenza (HITS)") & (d.anno == 2019)]
check("settori HITS 2019 tasso_natalita == 9.9", approx(row.iloc[0].tasso_natalita, 9.9))
row = d[(d.settore == "INDUSTRIA in senso stretto") & (d.anno == 2019)]
check("settori INDUSTRIA in s.s. 2019 imprese_nate == 18529", int(row.iloc[0].imprese_nate) == 18529)
check("settori: 13 settori", d.settore.nunique() == 13)
check("settori: 78 righe (13x6)", len(d) == 78)

# ed1924_regioni
d = load(f"{PREFIX}_regioni.csv")
check("regioni Campania 2019 tasso_natalita == 9.8",
      approx(d[(d.territorio == "Campania") & (d.anno == 2019)].iloc[0].tasso_natalita, 9.8))
check("regioni Piemonte 2019 tasso_natalita == 6.6",
      approx(d[(d.territorio == "Piemonte") & (d.anno == 2019)].iloc[0].tasso_natalita, 6.6))
check("regioni: 26 territori", d.territorio.nunique() == 26)
check("regioni: 156 righe (26x6)", len(d) == 156)

# ed1924_sopravvivenza
d = load(f"{PREFIX}_sopravvivenza.csv")
row = d[(d.macrosettore == "Industria in s.s.") & (d.coorte == 2019) & (d.anno_osservazione == 2024)]
check("sopravvivenza Industria in s.s. coorte2019 obs2024 == 55.307895731016245",
      approx(row.iloc[0].tasso_sopravvivenza, 55.307895731016245))
check("sopravvivenza: anni_dalla_nascita coerente",
      bool((d.anni_dalla_nascita == (d.anno_osservazione - d.coorte)).all()))
check("sopravvivenza: 5 macrosettori", d.macrosettore.nunique() == 5)
check("sopravvivenza: 75 righe (5 macro x 15 coppie triangolari)", len(d) == 75)

# ed1924_addetti_coorte2019
d = load(f"{PREFIX}_addetti_coorte2019.csv")
check("addetti: 5 macrosettori", d.macrosettore.nunique() == 5)
check("addetti Industria in s.s. addetti_t0_nate ~ 32198.56",
      approx(d[d.macrosettore == "Industria in senso stretto"].iloc[0].addetti_t0_nate, 32198.5600000008, tol=1e-6))

# ed1924_serie_tassi_2006_2024
d = load(f"{PREFIX}_serie_tassi_2006_2024.csv")
check("serie 2024 tasso_natalita == 7.3297644207251285",
      approx(d[d.anno == 2024].iloc[0].tasso_natalita, 7.3297644207251285))
check("serie 2024 tasso_mortalita == 6.400485267971362",
      approx(d[d.anno == 2024].iloc[0].tasso_mortalita, 6.400485267971362))
check("serie 2006 tasso_natalita == 7.13846075033129",
      approx(d[d.anno == 2006].iloc[0].tasso_natalita, 7.13846075033129))
check("serie: anni 2006..2024 (19)", len(d) == 19 and d.anno.min() == 2006 and d.anno.max() == 2024)
check("serie 2024 provvisorio == True", bool(d[d.anno == 2024].iloc[0].provvisorio) == True)

n_ok = sum(1 for _, ok in checks if ok)
print(f"\nCONTROLLI: {n_ok}/{len(checks)} OK")
if anomalies:
    print("\nANOMALIE:")
    for a in anomalies:
        print("  -", a)
else:
    print("\nNessuna anomalia.")

if n_ok != len(checks):
    raise SystemExit(f"VALIDAZIONE FALLITA: {len(checks) - n_ok} controlli KO")
print("\nTUTTI I CONTROLLI OK.")
