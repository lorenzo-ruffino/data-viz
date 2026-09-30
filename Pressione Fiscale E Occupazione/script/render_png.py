"""Rende ogni foglio di output/Tabelle_pressione_fiscale.xlsx in un PNG (output/screenshot/).
Legge valori e stili dal file Excel (unica fonte), valuta le formule, genera HTML e lo
fotografa con Chrome headless; poi ritaglia i margini bianchi."""
import os, re, subprocess, html
import openpyxl
from PIL import Image, ImageChops

BASE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "output")
XL = os.path.join(BASE, "Tabelle_pressione_fiscale.xlsx")
OUT = os.path.join(BASE, "screenshot"); os.makedirs(OUT, exist_ok=True)
CHROME = "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome"

wb = openpyxl.load_workbook(XL); cache = {}
def val(sh, ref):
    k = (sh, ref)
    if k in cache: return cache[k]
    v = wb[sh][ref].value
    if isinstance(v, str) and v.startswith("="):
        e = v[1:]
        e = re.sub(r"SUM\(([A-Z])(\d+):([A-Z])(\d+)\)", lambda m: "(" + "+".join(
            f"V('{sh}','{m.group(1)}{r}')" for r in range(int(m.group(2)), int(m.group(4)) + 1)) + ")", e)
        e = re.sub(r"'([^']+)'!([A-Z]\d+)", lambda m: f"V('{m.group(1)}','{m.group(2)}')", e)
        e = re.sub(r"(?<![A-Za-z'])([A-Z]\d+)(?![\d'])", lambda m: f"V('{sh}','{m.group(1)}')", e)
        v = eval(e, {"V": lambda s, r: (val(s, r) or 0)})
    cache[k] = v; return v

MINUS = "−"
def fmt(v, nf):
    if not isinstance(v, (int, float)): return html.escape(str(v))
    def it(x, d):  # formato italiano
        s = f"{abs(x):,.{d}f}".replace(",", "X").replace(".", ",").replace("X", ".")
        return s
    if nf.startswith("+0.00"):
        if abs(v) < 1e-9: return "\u2013"
        r = round(v, 2); return ("+" if r > 0 else MINUS if r < 0 else "") + it(v, 2)
    if nf.startswith("+0.0"):
        r = round(v, 1); return ("+" if r > 0 else MINUS if r < 0 else "") + it(v, 1)
    if nf == "0.0%": return (MINUS if v < 0 else "") + it(v * 100, 1) + "%"
    if nf == "0.00%": return (MINUS if v < 0 else "") + it(v * 100, 2) + "%"
    if nf == "#,##0": return it(v, 0)
    if nf == "0.0": return (MINUS if v < 0 else "") + it(v, 1)
    return it(v, 2)

CSS = """
@import url('https://fonts.googleapis.com/css2?family=Source+Sans+3:wght@400;600;700&display=swap');
body{margin:0;background:#fff;font-family:'Source Sans 3',Arial,sans-serif;color:#111827}
.card{display:inline-block;padding:28px 32px 22px 32px;background:#fff}
h1{font-size:25px;margin:0 0 4px 0;font-weight:700;color:#111827}
.sub{font-size:16px;color:#4b5563;margin:0 0 16px 0}
table{border-collapse:collapse;font-size:16px}
th{background:#1e3a8a;color:#fff;font-weight:600;padding:9px 14px;text-align:center}
th:first-child{text-align:left}
td{padding:6px 14px;border-bottom:1px solid #e5e7eb;text-align:center;white-space:nowrap}
td:first-child{text-align:left}
tr.sec td{background:#e0e7ff;color:#1e3a8a;font-weight:700;border-bottom:none}
tr.tot td{background:#f3f4f6;font-weight:700}
tr.topb td{border-top:2px solid #1e3a8a}
tr.muted td{color:#6b7280;font-style:italic;font-size:14px;border-bottom:none}
td.ind{padding-left:32px}
.cred{margin-top:14px;font-size:14px;font-weight:600;color:#374151}
.note{font-size:13px;color:#6b7280;margin-top:3px}
caption{caption-side:bottom;text-align:left}
"""

def sheet_html(name):
    ws = wb[name]
    title, sub = ws["B2"].value, ws["B3"].value
    rows_html, cred, notes = [], None, []
    maxc = ws.max_column
    for r in range(5, ws.max_row + 1):
        b = ws.cell(r, 2)
        cells = [ws.cell(r, c) for c in range(2, maxc + 1)]
        if all(c.value is None for c in cells): continue
        if r == 5:
            rows_html.append("<tr>" + "".join(f"<th>{html.escape(str(c.value or ''))}</th>" for c in cells if c.value is not None or c.column == 2) + "</tr>")
            ncols = sum(1 for c in cells if c.value is not None or c.column == 2); continue
        if b.font and b.font.italic or (b.font and b.font.bold and b.font.sz and b.font.sz <= 9):
            txt = str(b.value)
            if txt.startswith("Elaborazione"): cred = txt
            elif txt.startswith("Nota"): notes.append(txt)
            elif notes: notes[-1] += " " + txt
            if any(ws.cell(r, c).value is not None for c in range(3, maxc + 1)):  # riga "scarto": tabella
                pass
            else:
                continue
        fill = (b.fill.fgColor.rgb or "")[-6:] if b.fill and b.fill.fill_type else ""
        cls = []
        if fill == "E0E7FF": cls.append("sec")
        if fill == "F3F4F6": cls.append("tot")
        if b.border and b.border.top and b.border.top.style == "medium": cls.append("topb")
        if b.font and b.font.italic: cls.append("muted")
        ind = " class='ind'" if b.alignment and b.alignment.indent else ""
        tds = [f"<td{ind}>{html.escape(str(b.value))}</td>"]
        for c in range(3, 2 + ncols):
            cell = ws.cell(r, c)
            tds.append(f"<td>{fmt(val(name, cell.coordinate), cell.number_format)}</td>")
        rows_html.append(f"<tr class='{' '.join(cls)}'>" + "".join(tds) + "</tr>")
    body = (f"<div class='card'><h1>{html.escape(title)}</h1><div class='sub'>{html.escape(sub)}</div>"
            f"<table>{''.join(rows_html)}<caption><div class='cred'>{html.escape(cred or '')}</div>"
            + "".join(f"<div class='note'>{html.escape(n)}</div>" for n in notes) + "</caption></table></div>")
    return f"<!doctype html><html><head><meta charset='utf-8'><style>{CSS}</style></head><body>{body}</body></html>"

def shoot(htmlfile, png):
    subprocess.run([CHROME, "--headless=new", "--disable-gpu", "--hide-scrollbars", "--force-device-scale-factor=2",
                    "--window-size=1400,1300", "--virtual-time-budget=4000", f"--screenshot={png}", f"file://{htmlfile}"],
                   check=True, capture_output=True)
    im = Image.open(png).convert("RGB")
    bg = Image.new("RGB", im.size, (255, 255, 255))
    box = ImageChops.difference(im, bg).getbbox()
    pad = 40
    im.crop((max(box[0] - pad, 0), max(box[1] - pad, 0), min(box[2] + pad, im.width), min(box[3] + pad, im.height))).save(png)

for i, name in enumerate([s for s in wb.sheetnames if s[0].isdigit()], 1):
    slug = re.sub(r"[^a-z0-9]+", "_", name.lower()).strip("_")
    hf = os.path.join(OUT, f"{slug}.html"); pf = os.path.join(OUT, f"{slug}.png")
    open(hf, "w").write(sheet_html(name)); shoot(hf, pf); os.remove(hf)
    print(pf)
