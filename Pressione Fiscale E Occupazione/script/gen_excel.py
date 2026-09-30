"""Genera output/Tabelle_pressione_fiscale.xlsx: tabelle per screenshot (thread).
Valori calcolati da decomp_istat.py con le stesse ipotesi di dettaglio_intensita.py."""
import os
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.comments import Comment
import decomp_istat as D
from decomp_istat import S, x, step

OUT = os.path.join(os.path.dirname(os.path.abspath(__file__)), "..", "output")
os.makedirs(OUT, exist_ok=True)

# ---------------- calcoli (stesse ipotesi di dettaglio_intensita.py)
EPS, MARG = 1.7, 0.28
def calc(y0, y1):
    R, m = step(y0, y1); Y1 = S["PIL"][y1]
    g = lambda col, grp: R[R.gruppo.isin(grp)][col].sum()
    row = lambda v: float(R[R.voce == v].intens.iloc[0])
    occ_tot = m["L"] * ((S["ULA_TOT"][y1] / S["ULA_TOT"][y0]) / m["gYr"] - 1)
    comp = {
        "occ": occ_tot, "ricomp": m["occ"] - occ_tot, "sal": m["sal"],
        "prof": g("compos", ["A", "K"]), "altri": g("compos", ["P", "C", "E"]),
    }
    r0 = x("D613CE_T_C_W0", y0) / S["D11"][y0]; r1 = x("D613CE_T_C_W0", y1) / S["D11"][y1]
    dssc = (r1 - r0) * S["D11"][y1]; refl = -MARG * dssc
    qd, qp = {"2024": (0.533, 0.302), "2025": (0.539, 0.309)}[y1]
    drag = (qd * m["irp0"] * (EPS - 1) * (m["gW"] - 1) + qp * m["irp0"] * (EPS - 1) * (m["gPens"] - 1 - 0.003))
    detr = -8.44 if y1 == "2025" else 0.0
    rif = -4.3 if y1 == "2024" else 0.0
    irp_int = R[R.voce.str.startswith("IRPEF")].intens.sum()
    polizze = 100 * (x("D91C_C10_C_W0", y1) / Y1 - x("D91C_C10_C_W0", y0) / S["PIL"][y0])
    inten = {
        "fin": row("Redditi e rendite finanziarie") - polizze, "polizze": polizze,
        "cuneo": row("Contributi a carico dei dipendenti") + 100 * (refl + detr) / Y1,
        "riforma": 100 * rif / Y1, "drag": 100 * drag / Y1,
        "resid": irp_int - 100 * (refl + detr + rif + drag) / Y1,
        "energia": row("Energia (accise, oneri di sistema)"), "imprese": row("IRES, IRAP, sostitutive imprese"),
        "autonomi": row("Contributi autonomi + concordato"), "datori": row("Contributi a carico dei datori"),
        "iva": row("IVA"), "immobili": row("Immobili (IMU, cedolare, registro)"),
        "altro": row("Altro (tabacchi, giochi, bollo auto, canone, ETS, ...)"),
    }
    cuneo = dict(ssc=dssc, detr=detr, refl=refl, r0=r0, r1=r1)
    return comp, inten, cuneo, m
C24, I24, K24, M24 = calc("2023", "2024")
C25, I25, K25, M25 = calc("2024", "2025")

# ---------------- stile
ARIAL = "Arial"
F_TIT = Font(name=ARIAL, size=15, bold=True, color="1F2937")
F_SUB = Font(name=ARIAL, size=10.5, italic=False, color="4B5563")
F_HDR = Font(name=ARIAL, size=10.5, bold=True, color="FFFFFF")
F_TXT = Font(name=ARIAL, size=10.5, color="111827")
F_BOLD = Font(name=ARIAL, size=10.5, bold=True, color="111827")
F_SEC = Font(name=ARIAL, size=10.5, bold=True, color="1E3A8A")
F_SRC = Font(name=ARIAL, size=8.5, italic=True, color="6B7280")
F_CRED = Font(name=ARIAL, size=9, bold=True, color="374151")
FILL_HDR = PatternFill("solid", fgColor="1E3A8A")
FILL_SEC = PatternFill("solid", fgColor="E0E7FF")
FILL_TOT = PatternFill("solid", fgColor="F3F4F6")
THIN = Side(style="thin", color="9CA3AF"); THICK = Side(style="medium", color="1E3A8A")
PP = '+0.00;-0.00;0.00'          # punti di PIL
MLD = '0.0'                       # miliardi
PCT = '0.0%'
PCT2 = '0.00%'

def base(ws, title, subtitle, widths):
    ws.sheet_view.showGridLines = False
    for i, w in enumerate(widths): ws.column_dimensions[chr(66 + i)].width = w
    ws.column_dimensions["A"].width = 2
    ws["B2"] = title; ws["B2"].font = F_TIT
    ws["B3"] = subtitle; ws["B3"].font = F_SUB
    ws.row_dimensions[2].height = 22
def header(ws, r, labels):
    for j, l in enumerate(labels):
        c = ws.cell(row=r, column=2 + j, value=l); c.font = F_HDR; c.fill = FILL_HDR
        c.alignment = Alignment(horizontal="left" if j == 0 else "center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30
def line(ws, r, label, vals, fmt, font=F_TXT, fill=None, top=None, indent=0):
    c = ws.cell(row=r, column=2, value=label); c.font = font; c.alignment = Alignment(indent=indent, vertical="center")
    if fill: c.fill = fill
    if top: c.border = Border(top=top)
    for j, v in enumerate(vals):
        cc = ws.cell(row=r, column=3 + j, value=v); cc.font = font; cc.number_format = fmt
        cc.alignment = Alignment(horizontal="center", vertical="center")
        if fill: cc.fill = fill
        if top: cc.border = Border(top=top)
    ws.row_dimensions[r].height = 18
def source(ws, r, text, main=False):
    ws.cell(row=r, column=2, value=text).font = F_CRED if main else F_SRC
    ws.cell(row=r, column=2).alignment = Alignment(wrap_text=False)

wb = Workbook()
# ======================= 1. Pressione fiscale
ws = wb.active; ws.title = "1 Pressione"
base(ws, "La pressione fiscale torna a salire", "Pressione fiscale e peso dei redditi da lavoro dipendente sul PIL (%)", [44, 11, 11, 11, 11, 11])
years = ["2021", "2022", "2023", "2024", "2025"]
header(ws, 5, ["", *years])
press = [42.400, 41.642, 41.227, 42.212, 42.898]
pil = [1838674.8, 1998269.1, 2141760.8, 2210594.7, 2265002.5]
d1 = [738233.0, 783566.6, 823557.5, 868445.0, 902086.4]
line(ws, 6, "Pressione fiscale (% del PIL)", [p / 100 for p in press], PCT)
line(ws, 7, "Redditi da lavoro dipendente (milioni)", d1, '#,##0')
line(ws, 8, "PIL (milioni)", pil, '#,##0')
line(ws, 9, "Redditi da lavoro dipendente (% del PIL)", [f"={c}7/{c}8" for c in "CDEFG"], PCT, font=F_BOLD, fill=FILL_TOT, top=THIN)
ws["C6"].comment = Comment("Istat, Conti economici nazionali, 22 settembre 2026, Tavola 19", "LR")
source(ws, 11, "Elaborazione di Lorenzo Ruffino su dati Istat (Conti economici nazionali, 22 settembre 2026).", True)
source(ws, 12, "Nota: i dati 2024 e 2025 sono provvisori.")

# ======================= 2. Scomposizione
ws = wb.create_sheet("2 Scomposizione")
base(ws, "Perché è salita la pressione fiscale", "Variazione della pressione fiscale scomposta, in punti di PIL", [70, 13, 13])
header(ws, 5, ["", "2024", "2025"])
r = 6
line(ws, r, "Variazione della pressione fiscale (Istat)", [f"='1 Pressione'!F6*100-'1 Pressione'!E6*100", f"='1 Pressione'!G6*100-'1 Pressione'!F6*100"], PP, font=F_BOLD, fill=FILL_TOT); r += 2
def block(r, title, rows, src):
    r0 = r
    line(ws, r, title, [f"=SUM(C{r+1}:C{r+len(rows)})", f"=SUM(D{r+1}:D{r+len(rows)})"], PP, font=F_SEC, fill=FILL_SEC); r += 1
    for lab, k in rows:
        line(ws, r, lab, [src[0][k], src[1][k]], PP, indent=2); r += 1
    return r0, r + 1
comp_rows = [("Più lavoro rispetto al PIL", "occ"), ("Più dipendenti e meno autonomi", "ricomp"),
             ("Salari reali per unità di lavoro", "sal"), ("Minor peso di profitti e redditi degli autonomi", "prof"),
             ("Pensioni, consumi ed energia", "altri")]
ra, r = block(r, "A. Composizione dei redditi (le basi nel PIL crescono più o meno del PIL)", comp_rows, (C24, C25))
fin_rows = [("Interessi, dividendi, plusvalenze, risparmio gestito, bollo", "fin")]
rb, r = block(r, "B. Base finanziaria (redditi e patrimoni finanziari crescono più del PIL)", fin_rows, (I24, I25))
int_rows = [("Cuneo: decontribuzione (2024), detrazione e bonus (2025)", "cuneo"),
            ("Riforma IRPEF a tre aliquote", "riforma"), ("Fiscal drag IRPEF (stima)", "drag"),
            ("IRPEF, residuo (saldi e acconti, altro)", "resid"), ("Energia: fine degli sconti e oneri di sistema", "energia"),
            ("Prelievo straordinario sulle polizze vita (una tantum)", "polizze"),
            ("Imprese: IRES, IRAP, extraprofitti, rivalutazioni", "imprese"), ("Contributi degli autonomi", "autonomi"),
            ("Contributi a carico dei datori", "datori"), ("IVA", "iva"), ("Immobili: IMU, cedolare, registro", "immobili"),
            ("Altro: tabacchi, giochi, ETS, fondi di garanzia", "altro")]
rc, r = block(r, "C. Intensità (quanto si paga a parità di base: misure, fiscal drag, altro)", int_rows, (I24, I25))
line(ws, r, "A + B + C", [f"=C{ra}+C{rb}+C{rc}", f"=D{ra}+D{rb}+D{rc}"], PP, font=F_BOLD, fill=FILL_TOT, top=THICK); rt = r; r += 1
line(ws, r, "Scarto rispetto al dato Istat (arrotondamenti e revisioni)", [f"=C{rt}-C6", f"=D{rt}-D6"], PP, font=F_SRC); r += 2
source(ws, r, "Elaborazione di Lorenzo Ruffino su dati Istat (conti nazionali, settembre 2026; gettito per imposta, aprile 2026), MEF e UPB.", True); r += 1
source(ws, r, "Nota: eventuali mancate quadrature sono dovute agli arrotondamenti."); r += 1
source(ws, r, "Nota: in B le aliquote sui redditi finanziari non cambiano; il gettito cresce perché crescono interessi, plusvalenze e patrimoni, che non fanno parte del PIL."); r += 1
source(ws, r, "Nota: basi dello stesso anno; con le basi dell'anno prima per IRES, IRAP e autonomi, nel 2024 A sale e C scende di circa 0,6 punti (profitti 2023 incassati nel 2024)."); r += 1
source(ws, r, "Nota: fiscal drag stimato con elasticità IRPEF 1,7 (UPB); riflesso IRPEF della decontribuzione con aliquota marginale 28%.")

# ======================= 3. Occupazione
ws = wb.create_sheet("3 Occupazione")
base(ws, "Quanto conta l'occupazione", "Crescita del lavoro e suo contributo alla variazione della pressione fiscale", [62, 13, 13])
header(ws, 5, ["", "2024", "2025"])
ula_ind = {y: S["ULA_TOT"][y] - S["ULA_DIP"][y] for y in ["2023", "2024", "2025"]}
gr = lambda d: [d["2024"] / d["2023"] - 1, d["2025"] / d["2024"] - 1]
line(ws, 6, "Unità di lavoro dipendenti (var. %)", gr(S["ULA_DIP"]), PCT)
line(ws, 7, "Unità di lavoro indipendenti (var. %)", gr(ula_ind), PCT)
line(ws, 8, "Unità di lavoro totali (var. %)", gr(S["ULA_TOT"]), PCT)
line(ws, 9, "PIL reale (var. %)", gr(S["PILr"]), PCT)
line(ws, 11, "Variazione della pressione fiscale (punti di PIL)", ["='2 Scomposizione'!C6", "='2 Scomposizione'!D6"], PP, font=F_BOLD, fill=FILL_TOT)
line(ws, 12, "Più lavoro rispetto al PIL (unità di lavoro totali)", [C24["occ"], C25["occ"]], PP, indent=2)
line(ws, 13, "Più dipendenti e meno autonomi (a parità di lavoro totale)", [C24["ricomp"], C25["ricomp"]], PP, indent=2)
line(ws, 14, "Effetto dell'occupazione, totale", ["=C12+C13", "=D12+D13"], PP, font=F_BOLD, top=THIN)
line(ws, 16, "Quota dell'aumento: più lavoro rispetto al PIL", ["=C12/C11", "=D12/D11"], PCT)
line(ws, 17, "Quota dell'aumento: più dipendenti e meno autonomi", ["=C13/C11", "=D13/D11"], PCT)
line(ws, 18, "Quota dell'aumento: effetto dell'occupazione, totale", ["=C14/C11", "=D14/D11"], PCT, font=F_BOLD, fill=FILL_TOT, top=THIN)
source(ws, 20, "Elaborazione di Lorenzo Ruffino su dati Istat (Conti economici nazionali, 22 settembre 2026).", True)
source(ws, 21, "Nota: il lavoro dipendente è tassato più di quello autonomo: se una parte degli autonomi diventa dipendente,")
source(ws, 22, "il gettito sale anche senza nuovi posti di lavoro. Nel 2025 gli autonomi crescono più dei dipendenti e l'effetto si inverte.")

# ======================= 4. Redditi finanziari
ws = wb.create_sheet("4 Finanza")
base(ws, "Il fisco incassa di più dai redditi finanziari", "Gettito delle imposte su interessi, dividendi, plusvalenze e patrimoni finanziari (miliardi di euro)", [52, 11, 11, 11, 11])
header(ws, 5, ["", "2022", "2023", "2024", "2025"])
FIN = [("Ritenute sugli interessi (famiglie)", ["D51A_C04_C_W0"]), ("Dividendi (famiglie)", ["D51A_C10_C_W0"]),
       ("Risparmio gestito", ["D51C1_C02_C_W0"]), ("Plusvalenze su azioni", ["D51C1_C01_C_W0", "D51C2_C01_C_W0"]),
       ("Imposta di bollo", ["D214B_C04_C_W0"]), ("Assicurazioni vita e previdenza complementare", ["D51A_C12_C_W0", "D51A_C13_C_W0"]),
       ("Prelievo straordinario polizze vita", ["D91C_C10_C_W0"]),
       ("Interessi imprese, Tobin tax, cripto", ["D51B_C01_C_W0", "D214C_T_C_W0", "D51A_C15_C_W0"])]
import pandas as pd
dd = pd.read_csv(D.IN + "istat_95_815_DF_DCCN_FPA_5_2026M4.csv")
P = dd.pivot_table(index="DATA_TYPE_AGGR", columns=dd.TIME_PERIOD.astype(str), values="OBS_VALUE", aggfunc="sum") / 1000
fy = ["2022", "2023", "2024", "2025"]
r = 6
for lab, codes in FIN:
    line(ws, r, lab, [round(sum(float(P.loc[c, y]) if c in P.index else 0 for c in codes), 2) for y in fy], MLD); r += 1
line(ws, r, "Totale", [f"=SUM({c}6:{c}{r-1})" for c in "CDEF"], MLD, font=F_BOLD, fill=FILL_TOT, top=THICK); rt = r; r += 1
line(ws, r, "PIL (miliardi)", [1998.3, 2141.8, 2210.6, 2265.0], MLD); rp = r; r += 1
line(ws, r, "Totale in % del PIL", [f"={c}{rt}/{c}{rp}" for c in "CDEF"], PCT2, font=F_BOLD, fill=FILL_TOT); r += 2
source(ws, r, "Elaborazione di Lorenzo Ruffino su dati Istat (gettito per imposta, aprile 2026; PIL, 22 settembre 2026).", True); r += 1
source(ws, r, "Nota: l'imposta di bollo comprende anche quella su atti e documenti.")

# ======================= 5. Cuneo 2025
ws = wb.create_sheet("5 Cuneo 2025")
base(ws, "Il taglio del cuneo 2025 e la pressione fiscale", "Effetto sul gettito del passaggio da decontribuzione a detrazione più bonus (miliardi di euro)", [62, 14])
header(ws, 5, ["", "2025"])
line(ws, 6, "Contributi dei dipendenti in % delle retribuzioni, 2024", [K25["r0"]], PCT2)
line(ws, 7, "Contributi dei dipendenti in % delle retribuzioni, 2025", [K25["r1"]], PCT2)
line(ws, 9, "Contributi tornati pieni (fine della decontribuzione)", [round(K25["ssc"], 2)], '+0.0;-0.0')
line(ws, 10, "Nuova detrazione IRPEF (stima UPB)", [K25["detr"]], '+0.0;-0.0')
line(ws, 11, "Minore IRPEF perché i contributi riducono l'imponibile (stima)", [round(K25["refl"], 2)], '+0.0;-0.0')
line(ws, 12, "Effetto netto sulle entrate", ["=SUM(C9:C11)"], '+0.0;-0.0', font=F_BOLD, fill=FILL_TOT, top=THICK)
line(ws, 13, "Effetto netto in punti di PIL", ["=C12/2265.0*100"], PP, font=F_BOLD, fill=FILL_TOT)
line(ws, 15, "Per confronto: bonus fino a 20 mila euro, registrato come spesa (UPB)", [4.41], '0.0')
source(ws, 17, "Elaborazione di Lorenzo Ruffino su dati Istat e UPB (audizione sul DDL di bilancio 2025).", True)
source(ws, 18, "Nota: minore IRPEF stimata con un'aliquota marginale media del 28%.")

# ======================= 6. Contributi datori
ws = wb.create_sheet("6 Datori")
base(ws, "Contributi dei datori: crescono con i salari", "Contributi sociali a carico dei datori e retribuzioni lorde (miliardi di euro)", [48, 11, 11, 11])
header(ws, 5, ["", "2023", "2024", "2025"])
line(ws, 6, "Contributi a carico dei datori", [round(x("D611_T_C_W0", y), 1) for y in ["2023", "2024", "2025"]], MLD)
line(ws, 7, "Retribuzioni lorde", [round(S["D11"][y], 1) for y in ["2023", "2024", "2025"]], MLD)
line(ws, 8, "Contributi in % delle retribuzioni", ["=C6/C7", "=D6/D7", "=E6/E7"], PCT2, font=F_BOLD, fill=FILL_TOT, top=THIN)
source(ws, 10, "Elaborazione di Lorenzo Ruffino su dati Istat (retribuzioni: 22 settembre 2026; contributi: aprile 2026).", True)

# ======================= Note
ws = wb.create_sheet("Note e fonti")
ws.column_dimensions["B"].width = 130; ws.sheet_view.showGridLines = False
notes = [
 "Metodo: per ogni imposta, variazione del rapporto gettito/PIL = effetto base (la base cresce più/meno del PIL) + intensità (il gettito cresce più/meno della base).",
 "L'effetto base è diviso in A (basi che fanno parte del PIL: redditi da lavoro, profitti, consumi) e B (redditi e patrimoni finanziari, rapportati al PIL).",
 "Basi: retribuzioni lorde (contributi e IRPEF dei dipendenti), pensioni (IRPEF pensionati; prestazioni sociali in denaro al netto del bonus cuneo 2025),",
 "   risultato lordo di gestione + reddito misto (IRES, IRAP, autonomi), consumi sul territorio (IVA, energia), PIL (finanza, immobili, altro).",
 "IRPEF + addizionali ripartita con le quote MEF dell'imposta netta: dipendenti 53,3% (a.i. 2023) e 53,9% (a.i. 2024), pensionati 30,2% e 30,9%.",
 "Occupazione: prelievo sul lavoro dipendente × (crescita unità di lavoro / crescita PIL reale − 1).",
 "Misure: riforma IRPEF 2024 = 4,3 mld (MEF); cuneo 2025 = detrazione 8,44 mld + bonus 4,41 mld (UPB); decontribuzione 2024 = 10,7 mld netti (UPB).",
 "Robustezza: con base dell'anno prima per IRES/IRAP/autonomi, 2024 composizione +0,77 e intensità +0,21; fiscal drag con elasticità 1,5-1,7: 2024 +0,13/+0,18, 2025 +0,10/+0,14.",
 "Fonti: Istat, Conti economici nazionali 2010-2025 (22/9/2026); Istat esploradati 95_815_DF_DCCN_FPA_5 (ed. 2026M4); MEF, Bollettino entrate dicembre 2025;",
 "   MEF, Statistiche dichiarazioni IRPEF; UPB, Audizioni sui DDL di bilancio 2024 e 2025; Banca d'Italia, Relazione annuale sul 2025.",
]
ws["B2"] = "Note metodologiche e fonti"; ws["B2"].font = F_TIT
for i, n in enumerate(notes): ws.cell(row=4 + i, column=2, value=n).font = F_TXT

fn = os.path.join(OUT, "Tabelle_pressione_fiscale.xlsx")
wb.calculation.fullCalcOnLoad = True
wb.save(fn); print(fn)
